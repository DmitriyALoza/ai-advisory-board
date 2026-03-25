import io
import re
import subprocess
import tempfile
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader
from pptx import Presentation

from models.inputs import DocumentInput

MAX_CHARS_PER_DOCUMENT = 14000
MAX_TOTAL_CONTEXT_CHARS = 50000


def _clean_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def _extract_pdf(raw: bytes) -> str:
    reader = PdfReader(io.BytesIO(raw))
    parts = []
    for page in reader.pages:
        parts.append(page.extract_text() or "")
    return "\n".join(parts)


def _extract_docx(raw: bytes) -> str:
    doc = DocxDocument(io.BytesIO(raw))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n".join(parts)


def _extract_pptx(raw: bytes) -> str:
    prs = Presentation(io.BytesIO(raw))
    parts = []
    for slide_num, slide in enumerate(prs.slides, start=1):
        slide_parts = []
        for shape in slide.shapes:
            text = ""
            if hasattr(shape, "text"):
                text = shape.text or ""
            if text.strip():
                slide_parts.append(text.strip())
        if slide_parts:
            parts.append(f"Slide {slide_num}: " + " | ".join(slide_parts))
    return "\n".join(parts)


def _extract_text_file(raw: bytes) -> str:
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        return raw.decode("latin-1", errors="ignore")


def _extract_ppt_via_libreoffice(raw: bytes, filename: str) -> str:
    with tempfile.TemporaryDirectory() as tmp_dir:
        in_path = Path(tmp_dir) / filename
        in_path.write_bytes(raw)

        cmd = ["soffice", "--headless", "--convert-to", "pptx", "--outdir", tmp_dir, str(in_path)]
        result = subprocess.run(cmd, capture_output=True, text=True, check=False)
        if result.returncode != 0:
            err = result.stderr.strip() or result.stdout.strip() or "soffice conversion failed"
            raise RuntimeError(err)

        converted = Path(tmp_dir) / f"{in_path.stem}.pptx"
        if not converted.exists():
            raise RuntimeError("Converted .pptx file was not created")
        return _extract_pptx(converted.read_bytes())


def extract_documents(uploaded_files: list) -> tuple[list[DocumentInput], list[str]]:
    documents: list[DocumentInput] = []
    warnings: list[str] = []

    for file in uploaded_files:
        name = file.name
        suffix = Path(name).suffix.lower()
        raw = file.getvalue()

        try:
            if suffix == ".pdf":
                text = _extract_pdf(raw)
            elif suffix == ".docx":
                text = _extract_docx(raw)
            elif suffix == ".pptx":
                text = _extract_pptx(raw)
            elif suffix == ".ppt":
                text = _extract_ppt_via_libreoffice(raw, name)
            elif suffix in {".txt", ".md"}:
                text = _extract_text_file(raw)
            else:
                warnings.append(f"Skipped unsupported file type: {name}")
                continue
        except Exception as exc:
            warnings.append(f"Failed to parse {name}: {exc}")
            continue

        cleaned = _clean_text(text)
        if not cleaned:
            warnings.append(f"No readable text found in {name}")
            continue

        truncated = cleaned[:MAX_CHARS_PER_DOCUMENT]
        if len(cleaned) > MAX_CHARS_PER_DOCUMENT:
            warnings.append(f"Truncated long document for context window: {name}")

        documents.append(
            DocumentInput(
                filename=name,
                filetype=suffix.replace(".", ""),
                extracted_text=truncated,
            )
        )

    return documents, warnings


def build_document_digest(documents: list[DocumentInput]) -> str:
    if not documents:
        return "No documents uploaded."

    blocks = []
    used_chars = 0
    for doc in documents:
        header = f"Document: {doc.filename} ({doc.filetype})"
        budget_left = MAX_TOTAL_CONTEXT_CHARS - used_chars
        if budget_left <= 0:
            break

        body = doc.extracted_text[:budget_left]
        used_chars += len(body)
        blocks.append(f"{header}\n{body}")

    if not blocks:
        return "No usable document content available."

    return "\n\n".join(blocks)
